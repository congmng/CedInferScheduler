from docx import Document
import glob
for f in glob.glob('*.docx'):
    print('###', f)
    d=Document(f)
    for p in d.paragraphs:
        t=p.text.strip()
        if t: print(t)
    for ti,tab in enumerate(d.tables):
        print('TABLE',ti)
        for row in tab.rows:
            print(' | '.join(c.text.replace('\n',' / ') for c in row.cells))
