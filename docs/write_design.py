from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = '跨域PD分离存算协同调度具体设计与实验规划.docx'
d = Document()
sec=d.sections[0]
sec.top_margin=Cm(2.2); sec.bottom_margin=Cm(2.2); sec.left_margin=Cm(2.35); sec.right_margin=Cm(2.35)
styles=d.styles
styles['Normal'].font.name='Microsoft YaHei'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); styles['Normal'].font.size=Pt(10.5)
styles['Normal'].paragraph_format.space_after=Pt(6); styles['Normal'].paragraph_format.line_spacing=1.35
for name,size in [('Title',20),('Heading 1',15),('Heading 2',12),('Heading 3',11)]:
    s=styles[name]; s.font.name='Microsoft YaHei'; s._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); s.font.size=Pt(size); s.font.color.rgb=RGBColor(0,0,0)

def p(text='', style=None, boldlead=None):
    x=d.add_paragraph(style=style)
    if boldlead and text.startswith(boldlead):
        r=x.add_run(boldlead); r.bold=True; x.add_run(text[len(boldlead):])
    else: x.add_run(text)
    return x
def heading(text, level=1): d.add_heading(text, level=level)
def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)
def table(headers, rows, widths=None):
    t=d.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=h; shade(c,'1F4E78'); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in c.paragraphs[0].runs: r.font.color.rgb=RGBColor(255,255,255); r.bold=True; r.font.size=Pt(9)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,v in enumerate(row):
            cells[i].text=str(v); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri%2==1: shade(cells[i],'EAF2F8')
            for para in cells[i].paragraphs:
                para.paragraph_format.space_after=Pt(2); para.paragraph_format.space_before=Pt(2)
                for r in para.runs: r.font.size=Pt(9)
    if widths:
        for row in t.rows:
            for c,w in zip(row.cells,widths): c.width=Cm(w)
    d.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

d.add_paragraph('跨域 P D 分离推理中的', style='Title').alignment=WD_ALIGN_PARAGRAPH.CENTER
d.add_paragraph('存算协同结构重构调度 具体设计与实验规划', style='Title').alignment=WD_ALIGN_PARAGRAPH.CENTER
sub=d.add_paragraph('面向 Prefix 状态 异构 Decode 与跨域 P D 关系的轻量闭环方案'); sub.alignment=WD_ALIGN_PARAGRAPH.CENTER
sub.runs[0].italic=True

heading('一 结论与建议',1)
p('建议将课题收敛为“Cache aware Structural Gain Scheduling”。不把问题做成全局缓存放置或复杂强化学习，而是围绕一个可验证的核心闭环：先在当前资源结构下求最优 P-D 流量分配；当该分配仍暴露结构性失配时，再比较少数 P 实例增删候选带来的净收益；Prefix Cache 只通过改变 P 的有效供给和新实例的预热成本参与决策。这样既能明确体现存算协同，也能复用已有 P/D 分离、KV 传输和运行时路由能力。')
p('一句话问题：在跨域异构 P/D 分离推理中，如何利用 Prefix 状态校正 Prefill 的有效计算能力，并以“改变 P-D 关系矩阵后能否带来净收益”为依据，选择最小的 P 结构重构动作，从而降低 KV 传输与排队成本并提升 SLO goodput。')

heading('二 基于现有工作的定位',1)
table(['现有工作类型','可直接复用','尚未充分解决','本方案的切入点'],[
['P/D 弹性扩缩容','P/D 池、负载监测、按需求扩缩','实例通常被视为同质算力，按利用率或需求缺口扩缩','以 P-D 矩阵结构收益而非利用率单独触发 P 重构'],
['网络感知 P-D 路由','KV transfer、网络与队列感知 D 选择','多数只在固定资源池内选择 D','把 P 的新增、删除或位置变化视为改变可选 P-D 边的结构操作'],
['异构资源调度','GPU profiling、容量标定、放置候选','P 的真实供给受缓存状态影响，不能只由 GPU 型号决定','用 cache-aware effective capacity 描述 P 行能力'],
['Prefix Cache/复制','cache hit、热点统计、按需预热','缓存机制往往与扩缩容控制割裂','只在“扩某个 P 是否值得”时评估 cold/warm 两种新行'],
], [3.1,3.3,4.5,6.6])
p('与已有文档的关系：已有“P-D 关系矩阵驱动结构重构”已经给出最小成本流与 Structural Gain 主线；已有“跨域异构 P/D 分离”已覆盖缓存价值、异构 D、两时间尺度控制和较宽的联合优化。本设计选择前者作为主问题，将后者中最能体现存算协同、且工程上最轻的 cache-aware capacity 与 warm/cold 机制嵌入其中。')

heading('三 研究边界与创新点',1)
table(['保留','明确不做','理由'],[
['固定 D pool；P 侧每次仅允许 +1、-1 或 relocation','D 弹性、active Decode KV 迁移、权重迁移','避免离散决策爆炸，确保 P 行结构重构的因果可辨'],
['Prefix 状态统计、选择性热前缀预热','全局 Prefix placement/eviction 最优化','存储只服务于算力形成和重构价值判断'],
['min-cost flow/LP 作为内层 solver','复杂预测、RL、端到端 learned cost model','状态可由 EWMA 与 profiling 获得，方法可解释且易实现'],
['慢层重构 + 快层路由','新 request-level router','复用 serving runtime 的网络/队列感知选择'],
], [4.3,5.8,7.4])
p('创新点 1：结构收益而非容量阈值。P 扩缩容被形式化为 P-D 关系矩阵的一行增删；动作价值由重新求解后的全局分配改进减去重构成本决定。它可发现“总 P 容量充足但拓扑结构错误”的场景。')
p('创新点 2：状态相关的 P 行容量。P 的有效能力由 GPU 与缓存命中共同决定，而非实例 ready 即获得满额供给；这把 Prefix 存储状态直接接入计算调度。')
p('创新点 3：cold/warm 的同位候选对照。对同一扩容位置，分别评估冷启动与选择性预热，只有预热额外收益能覆盖复制成本时才触发。存储决策因此不是附加模块，而会改变动作排序。')

heading('四 系统模型与目标',1)
p('设 P={P1,...,Pm} 为 Prefill 实例集合，D={D1,...,Dn} 为固定 Decode 实例集合。每个 P_i 具有有效容量 μP_i，D_j 具有 profile 得到的解码容量 μD_j。A=[a_ij] 是可行矩阵；C=[c_ij] 是 pair cost；F=[f_ij] 是单位时间从 P_i 流向 D_j 的流量矩阵。')
p('Pair cost 采用可观测量：c_ij = RTT_ij + KVSize_i/BW_ij + Queue_j + DecodeTime_j。不同量纲可先归一化后加权，权重通过一个小验证集标定；第一版不学习该模型。')
p('给定资源结构 S，内层分配为：J*(S)=min_F Σ_iΣ_j f_ij c_ij；约束为 Σ_j f_ij≤μP_i，Σ_i f_ij≤μD_j，Σ_iΣ_j f_ij≥λ，f_ij=0 当 a_ij=0。该问题是容量约束最小成本流，可直接使用成熟求解器。')
p('缓存感知容量采用轻量估计：μP_i = ProfileP(GPU_i, effective_tokens_i)，其中 effective_tokens_i = ISL - HitPrefixLen_i。ProfileP 由离线 profile 表插值得到，在线仅用滑窗/EWMA 统计命中前缀长度与输入长度。')

heading('五 具体调度算法',1)
p('1. 观测。控制器每个慢周期收集 λ、ISL/OSL、P 侧 hit prefix length、P/D 队列、GPU profile、RTT/BW 与实例状态。')
p('2. 当前结构最优。构造 A、C、μP、μD，解出 F*(S) 和 J*(S)。运行时按 F 的 affinity 比例给请求选择候选 D，并在候选中用实时 queue/cost 打破平局。')
p('3. 轻量诊断。计算 assignment regret R=Σ_i λ_i(cbar_i-cbest_i)、高成本边占比 E_exp 与 D pressure ρ_j。它们只用于判断是否进入候选评估，不参与复杂的加权总目标。')
p('4. 反事实候选。若 SLO 恶化或诊断触发，仅构造邻域 N(S)：各可用域的 +1 P(cold)、+1 P(warm)，以及满足最小存活时间的 -1 P_i。warm 行仅复制 Top-K 热前缀，容量更高但有复制与等待成本。')
p('5. 选择动作。对每个 a∈N(S)，构造 S_a 并重求 J*(S_a)，计算 G(a)=J*(S)-J*(S_a)-C_reconfig(a)。选择 G 最大者；仅当 G(a)>θ、SLO 约束可行并满足 dwell time 时执行，否则维持结构。')
p('6. 执行。扩容先拉起 P；warm 候选并行复制 Top-K；缩容停止接新 Prefill，等待本地任务完成后释放。动作完成后更新 affinity。')
table(['候选','P 行变化','容量与成本','适用判断'],[
['+P@domain k cold','新增一行','μcold；启动成本 + 后续更多重算','热点短暂、复制带宽差或预期收益不够'],
['+P@domain k warm','新增一行','μwarm>μcold；启动 + Top-K 复制/等待','热点稳定且 Gwarm 超过 Gcold'],
['-P_i','删除一行','释放资源；cache loss/排空成本','删除后 J 上升最小且资源节约足够'],
['relocate P','删旧行加新行','改变 pair-cost vector','用于显示位置/网络而非数量造成的结构差异'],
], [3.2,3.0,5.3,6.0])

heading('六 存算协同如何被可验证地体现',1)
table(['链路','机制','可观测证据'],[
['Storage to Compute','Prefix hit 改变 effective tokens，进而改变 μP_i','同 GPU、不同 cache state 的 μP；Effective Capacity Ramp-up Time'],
['Storage to Structure','warm/cold 改变新增 P 行容量与重构成本，因此改变 G(a) 排序','记录 Gcold、Gwarm、选择结果与复制流量'],
['Compute to Storage','只有选中 warm 扩容时才复制 Top-K 高价值 prefix','复制集合、Saved Prefill GPU time、recomputed prefix tokens'],
], [3.3,7.0,7.2])
p('前缀价值可用 V_k = predicted_reuse_k × saved_prefill_time_k - copy_time_k 表示；工程上选 V_k>0 的 Top-K 即可。它不是完整 cache placement 优化，而是“给定一次 warm 动作，哪些状态值得随新 P 行形成”的小决策。')

heading('七 实现路线与整体流程',1)
p('建议分为“真实小规模闭环 + trace-driven 扩展”两阶段。真实系统负责证明 P/D、KV transfer、缓存与控制器能真实协同；仿真器负责系统性扫网络、异构度和工作负载。两者共享同一 profile 表、成本模型、控制算法和 trace 格式。')
table(['阶段','工作内容','产出/验收'],[
['M1 基线与测量','选定现有 P/D serving runtime；记录 KV 传输、cache hit、P/D queue；完成 GPU/ISL/OSL profile','profile 表；可复现 static matching'],
['M2 内层分配','实现状态采集、A/C 构造与 min-cost flow；将结果下发为 affinity','固定结构下优于 nearest-D/简单负载均衡'],
['M3 结构重构','实现诊断、±1 P counterfactual、G 阈值、dwell time','结构瓶颈场景中少量重构能降低 J 与 P95 TTFT'],
['M4 存算协同','实现 cache-aware μP、Top-K warm 与 cold/warm 评估','warm 是否执行可由净收益解释，缩短 ramp-up'],
['M5 完整评估','接入 trace、故障/带宽变化、基线与消融','完整曲线、统计显著性和开销报告'],
], [2.5,7.0,7.8])

heading('八 实验设计',1)
heading('8.1 环境与工作负载',2)
p('第一版仅需两域 Edge 与 Cloud：域内高带宽，域间可控 RTT/BW。D 固定且至少有两档 decode throughput（可由真实异构 GPU 或限速/profile 仿真形成）；每个域预留可启停 P 池。选 7B/14B 级模型，使用支持 P/D 分离与 Prefix Cache 的开源 serving stack。若真实资源不够，保留 2 域原型，使用 trace-driven 模拟器扩展至 3 域和更多实例。')
table(['Case','构造','要验证的结论'],[
['A 固定结构配对','固定 P/D，扫描 WAN、KV size、D capacity','显式 pair cost + 列容量优于 nearest/fastest-D'],
['B 容量足够但结构失配','总 μP≥λ，多个 P 仍被迫走高成本 D 边','容量未超载并不意味着不需要结构性扩容'],
['C 容量驱动对结构驱动','阶梯 λ 与 topology mismatch 混合','Ours 仅在 G 足够大时重构，GPU-hours 与 SLO 更优'],
['D warm 对 cold','固定扩容位置，扫描 prefix reuse 与复制带宽','缓存状态会改变候选排序与 ramp-up'],
['E scale-in','负载下降且 P 行缓存价值/拓扑不同','删除结构价值最低行优于按利用率删除'],
['F 热点漂移','λ 稳定但热门 prefix 切换','仅看负载的计算调度会误判有效容量'],
], [3.3,6.6,7.4])
heading('8.2 对照与消融',2)
table(['方法','共同条件','差异与目的'],[
['B0 Static + optimal matching','同一 cost matrix、min-cost flow、cache runtime','不改变 P 结构，给出固定结构上限'],
['B1 Load scaling + optimal matching','同上','按 utilization/demand gap ±P，检验结构收益触发'],
['B2 Pair-cost heuristic','同上','在高成本流量附近扩 P，但不做反事实重求'],
['B3 Structure-aware NoCache','同上','使用 G，但 μP 仅由 GPU 决定且无 warm/cold，隔离存储贡献'],
['Ours','同上','G + cache-aware μP + warm/cold'],
['消融 NoWarm / NoHysteresis','同上','验证预热收益与抗抖动机制'],
], [3.3,5.4,8.6])
p('公平性原则：所有方法都启用相同的 Prefix Cache runtime 与相同 P/D matching solver；差异只能来自“是否把 cache state 用于调度决策”。不能用简单的 Cache ON vs OFF 证明存算协同。')
heading('8.3 指标与报告方式',2)
table(['类别','核心指标','对应主张'],[
['服务性能','P50/P95/P99 TTFT、TPOT、SLO attainment、goodput','端到端价值'],
['分配质量','J*(S)、平均 pair cost、regret、高成本边占比、D load balance','矩阵优化和结构失配'],
['资源与网络','GPU-hours、活跃 P 数、P→D KV traffic、跨域峰值流量','效率与通信权衡'],
['存算协同','effective μP、recomputed prefix tokens、Saved Prefill GPU time、warmup traffic、ramp-up time','缓存真正改变计算与重构'],
['稳定性与开销','重构次数、配置翻转、planner latency、预测误差敏感性','工程可用性'],
], [3.0,7.6,6.7])
p('建议每个动态 case 至少运行 5 次或使用多段 trace，报告均值与 95% 置信区间；同时给出动作时间线（λ、J、活跃 P、warm/cold、SLO）以展示为何重构，而非只汇总平均指标。')

heading('九 预期论证链与风险控制',1)
table(['待证明假设','最直接的实验','失败时的解释与调整'],[
['H1 固定结构的 P-D 全局分配有价值','Case A','若提升弱，检查网络差异/D 列约束是否足够明显'],
['H2 容量充足仍会结构失配','Case B','人为构造高成本边与受限列，避免只在过载时比较'],
['H3 G 优于利用率扩缩','Case C/E','确保基线共享 matcher，收益归因于结构动作判断'],
['H4 cache state 改变结构选择','Case D/F、B3 对照','记录 Gwarm/Gcold 和 capacity ramp-up，而非只报告 hit rate'],
['H5 控制器可在线运行','planner latency 与候选数 sweep','保持候选为 O(域数+P数)，必要时仅触发后评估'],
], [4.2,6.3,6.8])
p('主要风险是把“缓存命中提升”误当作存算协同贡献。规避方法是固定缓存机制，并以 NoCache-aware controller、cold/warm counterfactual、recomputed tokens 与有效容量爬升时间证明缓存状态确实改变了重构决策。另一风险是成本模型不准；第一版可将其作为可替换模块，并报告通过测量值与预测值的偏差。')

heading('十 最小可发表闭环',1)
p('最小闭环只需四个模块：状态采集与 profile、min-cost-flow P-D 分配、Structural Gain 的 ±1 P 反事实评估、cache-aware cold/warm 新行。它能形成完整论文叙事：当前最优分配仍可能被资源结构锁死；新增/删除一行的价值来自重塑全局 P-D 分配；Prefix 状态既改变已有 P 的有效能力，也改变新行的形成成本，故存储状态必须进入计算重构决策。')
p('建议题目：面向跨域异构 P/D 分离推理的缓存感知结构收益调度。')

d.save(OUT)
print(OUT)
